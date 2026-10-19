from __future__ import annotations

from pathlib import Path
import sys

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from student_analytics.analysis import cluster_students, train_risk_models
from student_analytics.data import load_sample_data

ROOT = Path(__file__).resolve().parents[1]
SCREENSHOT = ROOT / "assets" / "dashboard_screenshot.png"
OUTPUT = ROOT / "Student_Academic_Performance_Analysis.pptx"

NAVY = "102B4E"
BLUE = "3B82D0"
GREEN = "28A36A"
TEAL = "39A99B"
ORANGE = "E99540"
RED = "D94F54"
INK = "182230"
MUTED = "667085"
PALE = "F5F7FB"
WHITE = "FFFFFF"
LINE = "E4E9F0"
FONT = "Aptos"
SW = 13.333
SH = 7.5


def color(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: str,
    *,
    rounded: bool = False,
    line: str | None = None,
):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.fill.background() if line is None else None
    if line is not None:
        shape.line.color.rgb = color(line)
        shape.line.width = Pt(1)
    if rounded:
        shape.adjustments[0] = 0.12
    return shape


def textbox(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 16,
    fill: str = INK,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0)
    frame.margin_right = Inches(0)
    frame.margin_top = Inches(0)
    frame.margin_bottom = Inches(0)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color(fill)
    return shape


def add_background(slide, fill: str = PALE) -> None:
    rect(slide, 0, 0, SW, SH, fill)


def add_header(slide, eyebrow: str, title: str, subtitle: str = "") -> None:
    textbox(slide, eyebrow.upper(), 0.68, 0.38, 12, 0.22, size=9, fill=BLUE, bold=True)
    textbox(slide, title, 0.68, 0.7, 12, 0.55, size=27, fill=NAVY, bold=True)
    if subtitle:
        textbox(slide, subtitle, 0.68, 1.32, 12, 0.38, size=11, fill=MUTED)


def add_footer(slide, number: int, dark: bool = False) -> None:
    shade = "C5D2E2" if dark else "98A2B3"
    textbox(slide, "STUDENT ACADEMIC PERFORMANCE ANALYSIS  ·  DATA MINING", 0.68, 7.17, 10.5, 0.16, size=7.5, fill=shade)
    textbox(slide, f"{number:02d}", 12.15, 7.13, 0.5, 0.2, size=9, fill=shade, align=PP_ALIGN.RIGHT)


def add_card(slide, x: float, y: float, w: float, h: float, title: str, body: str, accent: str = BLUE) -> None:
    rect(slide, x, y, w, h, WHITE, rounded=True, line=LINE)
    rect(slide, x, y, 0.07, h, accent)
    textbox(slide, title, x + 0.24, y + 0.2, w - 0.45, 0.34, size=15, fill=NAVY, bold=True)
    textbox(slide, body, x + 0.24, y + 0.66, w - 0.46, h - 0.79, size=11, fill=MUTED)


def add_kpi(slide, x: float, y: float, w: float, label: str, value: str, note: str, accent: str) -> None:
    rect(slide, x, y, w, 1.15, WHITE, rounded=True, line=LINE)
    textbox(slide, label, x + 0.18, y + 0.15, w - 0.35, 0.22, size=10, fill=MUTED, bold=True)
    textbox(slide, value, x + 0.18, y + 0.39, w - 0.35, 0.46, size=25, fill=accent, bold=True)
    textbox(slide, note, x + 0.18, y + 0.88, w - 0.35, 0.16, size=8.5, fill=MUTED)


def add_dot_bullet(slide, text: str, x: float, y: float, w: float, *, accent: str = BLUE, size: float = 11) -> None:
    rect(slide, x, y + 0.09, 0.09, 0.09, accent, rounded=True)
    textbox(slide, text, x + 0.2, y, w - 0.2, 0.45, size=size, fill=INK)


def build_deck() -> Path:
    data = load_sample_data()
    model_results = train_risk_models(data)
    if model_results is None:
        raise RuntimeError("Could not train the risk models on the bundled sample data.")
    segmented, silhouette = cluster_students(data)

    if not SCREENSHOT.is_file():
        raise FileNotFoundError(f"Dashboard screenshot not found: {SCREENSHOT}")

    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    blank = prs.slide_layouts[6]

    # 1. Cover
    slide = prs.slides.add_slide(blank)
    add_background(slide, NAVY)
    rect(slide, 9.1, 0, 4.233, SH, "15395F")
    for idx, (label, value, accent) in enumerate(
        [("PASS RATE", f"{data['pass_status'].eq('Pass').mean():.0%}", GREEN),
         ("ATTENDANCE", f"{data['attendance_rate'].mean():.0f}%", BLUE),
         ("STUDENTS", f"{len(data):,}", ORANGE)]
    ):
        y = 1.55 + idx * 1.48
        rect(slide, 9.72, y, 2.84, 1.12, "1D466D", rounded=True)
        textbox(slide, label, 9.98, y + 0.18, 2.3, 0.18, size=9, fill="C5D2E2", bold=True)
        textbox(slide, value, 9.98, y + 0.45, 2.3, 0.44, size=25, fill=accent, bold=True)
    rect(slide, 0.72, 0.72, 0.11, 0.55, TEAL, rounded=True)
    textbox(slide, "DATA MINING PROJECT", 1.0, 0.83, 6.2, 0.2, size=10, fill="AFC6DD", bold=True)
    textbox(slide, "Student Academic\nPerformance Analysis", 0.72, 1.75, 7.8, 1.65, size=34, fill=WHITE, bold=True)
    textbox(slide, "From student records to actionable academic insight", 0.75, 3.72, 7.4, 0.48, size=17, fill="DAE6F2")
    rect(slide, 0.75, 4.62, 2.35, 0.47, "244B72", rounded=True)
    textbox(slide, "PYTHON  ·  STREAMLIT  ·  SCIKIT-LEARN", 0.91, 4.77, 2.04, 0.14, size=8, fill=WHITE, bold=True)
    textbox(slide, "Descriptive analytics  +  predictive modeling  +  student segmentation", 0.75, 5.42, 7.7, 0.42, size=13, fill="AFC6DD")
    textbox(slide, "Synthetic sample data · 480 student records", 0.75, 6.55, 7.6, 0.26, size=10, fill="AFC6DD")
    add_footer(slide, 1, dark=True)

    # 2. Project overview
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_header(slide, "01 · Project overview", "One dashboard. Three ways to understand performance.",
               "A practical education analytics workflow designed for exploration and early support.")
    cards = [
        ("01  Explore outcomes", "Track pass rate, GPA estimate, attendance, score distributions, school and grade patterns.", BLUE),
        ("02  Identify risk", "Compare interpretable baseline classification models using engagement and prior-performance signals.", RED),
        ("03  Find segments", "Use unsupervised clustering to group students by learning and engagement patterns.", GREEN),
    ]
    for i, (title, body, accent) in enumerate(cards):
        add_card(slide, 0.7 + i * 4.18, 2.05, 3.86, 1.62, title, body, accent)
    kpis = [
        ("PASS RATE", f"{data['pass_status'].eq('Pass').mean():.0%}", "Score of 60 or above", GREEN),
        ("AVERAGE GPA", f"{min(4, data['final_score'].mean() / 25):.2f}", "Score mapped to 4.0 scale", BLUE),
        ("ATTENDANCE", f"{data['attendance_rate'].mean():.0f}%", "Mean attendance rate", TEAL),
        ("AT-RISK", f"{int(data['at_risk'].sum())}", "Score <60 or attendance <70%", RED),
    ]
    for i, item in enumerate(kpis):
        add_kpi(slide, 0.7 + i * 3.15, 4.15, 2.9, *item)
    rect(slide, 0.7, 5.75, 11.98, 0.72, "EAF2FB", rounded=True)
    textbox(slide, "Purpose", 0.94, 5.92, 1.0, 0.2, size=11, fill=NAVY, bold=True)
    textbox(slide, "Surface patterns that help educators investigate and support students—not automate decisions about them.",
            2.0, 5.88, 10.3, 0.28, size=12, fill=INK)
    add_footer(slide, 2)

    # 3. Data
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_header(slide, "02 · Dataset", "A small, reproducible dataset with meaningful learning signals.",
               "Synthetic records let the full pipeline run out of the box; replace them with validated data for further study.")
    add_kpi(slide, 0.7, 1.95, 2.82, "STUDENT RECORDS", f"{len(data):,}", "Generated sample rows", BLUE)
    add_kpi(slide, 3.68, 1.95, 2.82, "INPUT FIELDS", "13", "Required CSV columns", TEAL)
    add_kpi(slide, 6.66, 1.95, 2.82, "GRADE RANGE", "9–12", "High-school sample", GREEN)
    add_kpi(slide, 9.64, 1.95, 2.82, "ACADEMIC YEAR", "2024", "Assessment records", ORANGE)
    add_card(slide, 0.7, 3.45, 3.85, 2.38, "Academic history",
             "Previous and final scores\nGrade level\nAssessment date\nEnrollment year", BLUE)
    add_card(slide, 4.75, 3.45, 3.85, 2.38, "Engagement",
             "Attendance percentage\nAssignment completion\nStudy hours per week", GREEN)
    add_card(slide, 8.8, 3.45, 3.85, 2.38, "Learning context",
             "School\nParent / guardian education\nInternet access\nGender category", ORANGE)
    textbox(slide, "Derived outcome labels", 0.78, 6.12, 2.0, 0.2, size=10, fill=NAVY, bold=True)
    textbox(slide, "Pass = final score ≥ 60     ·     At-risk = final score < 60 OR attendance < 70%",
            2.65, 6.11, 9.8, 0.25, size=11, fill=MUTED)
    add_footer(slide, 3)

    # 4. Data mining workflow
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_header(slide, "03 · Method", "A transparent pipeline from raw CSV to useful views.",
               "Validation, descriptive analysis, supervised prediction, and unsupervised segmentation.")
    steps = [
        ("1", "Validate", "Required fields\nRanges · dates", BLUE),
        ("2", "Explore", "Score · attendance\nSchool · grade", TEAL),
        ("3", "Predict", "Logistic regression\nRandom forest", RED),
        ("4", "Segment", "Scaled features\nK-means", ORANGE),
        ("5", "Review", "Metrics · profiles\nEducator context", GREEN),
    ]
    for i, (number, title, body, accent) in enumerate(steps):
        x = 0.72 + i * 2.53
        rect(slide, x, 2.25, 2.12, 2.05, WHITE, rounded=True, line=LINE)
        rect(slide, x + 0.2, 2.48, 0.48, 0.48, accent, rounded=True)
        textbox(slide, number, x + 0.2, 2.59, 0.48, 0.18, size=12, fill=WHITE, bold=True, align=PP_ALIGN.CENTER)
        textbox(slide, title, x + 0.2, 3.18, 1.75, 0.28, size=16, fill=NAVY, bold=True)
        textbox(slide, body, x + 0.2, 3.58, 1.75, 0.55, size=10, fill=MUTED)
        if i < len(steps) - 1:
            rect(slide, x + 2.14, 3.18, 0.33, 0.04, "B8C5D5")
    rect(slide, 0.72, 4.8, 11.9, 1.2, "EAF2FB", rounded=True)
    textbox(slide, "Leakage guard", 0.98, 5.02, 1.58, 0.25, size=12, fill=NAVY, bold=True)
    textbox(slide, "Risk classification excludes final score and the derived risk label from model inputs. "
            "A stratified 75/25 split keeps test examples separate from training.",
            2.8, 4.98, 9.3, 0.55, size=11, fill=INK)
    add_footer(slide, 4)

    # 5. Actual dashboard screenshot
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    textbox(slide, "PRODUCT SCREENSHOT  ·  STREAMLIT OVERVIEW", 0.68, 0.32, 12, 0.22, size=9, fill=BLUE, bold=True)
    textbox(slide, "The dashboard at a glance", 0.68, 0.6, 12, 0.42, size=23, fill=NAVY, bold=True)
    slide.shapes.add_picture(str(SCREENSHOT), Inches(1.29), Inches(1.08), width=Inches(10.75), height=Inches(6.05))
    add_footer(slide, 5)

    # 6. Predictive risk analysis
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_header(slide, "04 · Predictive analysis", "Early-warning classification: measure more than accuracy.",
               "Two baseline classifiers, one stratified holdout, and clear performance trade-offs.")
    best = model_results["best_model"]
    best_metrics = model_results["metrics"][best]
    metric_items = [
        ("BEST MODEL", best, BLUE, "Selected by F1"),
        ("ACCURACY", f"{best_metrics['accuracy']:.1%}", GREEN, "Correct overall"),
        ("PRECISION", f"{best_metrics['precision']:.1%}", ORANGE, "Risk flags that were correct"),
        ("RECALL", f"{best_metrics['recall']:.1%}", RED, "At-risk cases identified"),
    ]
    for i, (label, value, accent, note) in enumerate(metric_items):
        add_kpi(slide, 0.7 + i * 3.15, 1.88, 2.9, label, value, note, accent)
    rect(slide, 0.7, 3.35, 5.7, 2.55, WHITE, rounded=True, line=LINE)
    textbox(slide, "Holdout comparison", 0.96, 3.59, 4.9, 0.3, size=15, fill=NAVY, bold=True)
    comparison = model_results["metrics"]
    metric_names = [("Accuracy", "accuracy"), ("Precision", "precision"), ("Recall", "recall"), ("F1", "f1")]
    for row, (label, key) in enumerate(metric_names):
        y = 4.08 + row * 0.4
        textbox(slide, label, 0.98, y, 1.1, 0.2, size=9.5, fill=MUTED)
        for col, model_name in enumerate(("Logistic Regression", "Random Forest")):
            value = comparison[model_name][key]
            x = 2.0 + col * 2.1
            rect(slide, x, y + 0.025, 1.38, 0.14, "E8EDF3", rounded=True)
            rect(slide, x, y + 0.025, max(0.05, 1.38 * value), 0.14, BLUE if col == 0 else GREEN, rounded=True)
            textbox(slide, f"{value:.0%}", x + 1.42, y - 0.01, 0.5, 0.2, size=9, fill=INK, bold=True)
    textbox(slide, "Logistic regression", 2.0, 5.7, 1.85, 0.18, size=8.3, fill=MUTED)
    textbox(slide, "Random forest", 4.1, 5.7, 1.5, 0.18, size=8.3, fill=MUTED)
    rect(slide, 6.68, 3.35, 5.95, 2.55, WHITE, rounded=True, line=LINE)
    textbox(slide, "Top random-forest drivers", 6.94, 3.59, 5.3, 0.3, size=15, fill=NAVY, bold=True)
    drivers = model_results["feature_importance"].head(5).sort_values("importance", ascending=True)
    max_importance = float(drivers["importance"].max())
    for row, (_, item) in enumerate(drivers.iterrows()):
        y = 4.04 + row * 0.34
        textbox(slide, str(item["feature"]), 6.96, y, 2.12, 0.2, size=9, fill=INK)
        rect(slide, 9.17, y + 0.025, 2.42, 0.13, "E8EDF3", rounded=True)
        rect(slide, 9.17, y + 0.025, 2.42 * float(item["importance"]) / max_importance, 0.13, TEAL, rounded=True)
    rect(slide, 0.7, 6.22, 11.93, 0.53, "FFF4E5", rounded=True)
    textbox(slide, "Interpretation", 0.93, 6.38, 1.3, 0.18, size=10, fill=NAVY, bold=True)
    textbox(slide, "Good accuracy can coexist with missed at-risk students; review recall, subgroup behavior, and context before use.",
            2.28, 6.34, 10.0, 0.22, size=10, fill=INK)
    add_footer(slide, 6)

    # 7. Segmentation
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_header(slide, "05 · Unsupervised analysis", "Student segments reveal different engagement patterns.",
               "K-means groups students using standardized attendance, assignment completion, study time, and prior score.")
    if silhouette is None:
        raise RuntimeError("Silhouette score could not be calculated for the sample segmentation.")
    add_kpi(slide, 0.7, 1.88, 2.72, "SEGMENTS", str(segmented["segment"].nunique()), "K-means groups", BLUE)
    add_kpi(slide, 3.56, 1.88, 2.72, "SILHOUETTE", f"{silhouette:.2f}", "Modest separation in this sample", ORANGE)
    add_kpi(slide, 6.42, 1.88, 2.72, "FEATURES", "4", "Standardized engagement inputs", TEAL)
    add_kpi(slide, 9.28, 1.88, 2.72, "RISK FLAG", "Derived", "Score and attendance rules", RED)

    profile = segmented.groupby("segment", as_index=False).agg(
        students=("student_id", "count"),
        attendance=("attendance_rate", "mean"),
        assignments=("assignment_completion", "mean"),
        study_hours=("study_hours", "mean"),
        at_risk=("at_risk", "mean"),
    ).sort_values("segment")
    rect(slide, 0.7, 3.38, 7.75, 2.65, WHITE, rounded=True, line=LINE)
    textbox(slide, "Segment profile", 0.96, 3.62, 5, 0.28, size=15, fill=NAVY, bold=True)
    headers = [("Segment", 0.98, 1.2), ("Students", 2.42, 0.85), ("Attend.", 3.48, 0.8),
               ("Assignments", 4.53, 1.05), ("Study hrs", 5.82, 0.95), ("At-risk", 7.0, 0.9)]
    for label, x, w in headers:
        textbox(slide, label, x, 4.08, w, 0.2, size=8.5, fill=MUTED, bold=True)
    for row, (_, item) in enumerate(profile.iterrows()):
        y = 4.46 + row * 0.34
        textbox(slide, str(item["segment"]), 0.98, y, 1.2, 0.2, size=9.5, fill=INK, bold=True)
        textbox(slide, f"{int(item['students'])}", 2.42, y, 0.85, 0.2, size=9.5, fill=INK)
        textbox(slide, f"{item['attendance']:.0f}%", 3.48, y, 0.8, 0.2, size=9.5, fill=INK)
        textbox(slide, f"{item['assignments']:.0f}%", 4.53, y, 1.05, 0.2, size=9.5, fill=INK)
        textbox(slide, f"{item['study_hours']:.1f}", 5.82, y, 0.95, 0.2, size=9.5, fill=INK)
        textbox(slide, f"{item['at_risk']:.0%}", 7.0, y, 0.9, 0.2, size=9.5, fill=RED if item["at_risk"] >= .2 else GREEN)

    rect(slide, 8.72, 3.38, 3.9, 2.65, "EAF2FB", rounded=True)
    textbox(slide, "What stands out", 8.99, 3.65, 3.25, 0.28, size=15, fill=NAVY, bold=True)
    riskiest = profile.sort_values("at_risk", ascending=False).iloc[0]
    add_dot_bullet(slide, f"{riskiest['segment']} has the highest observed risk rate ({riskiest['at_risk']:.0%}).",
                   9.0, 4.15, 3.2, accent=RED, size=10)
    add_dot_bullet(slide, "Clusters are exploratory—not fixed student types or diagnoses.",
                   9.0, 4.9, 3.2, accent=BLUE, size=10)
    textbox(slide, "Silhouette is modest: consider other features and cluster counts before interpreting groups.",
            0.82, 6.35, 11.6, 0.3, size=10, fill=MUTED)
    add_footer(slide, 7)

    # 8. Architecture and setup
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_header(slide, "06 · Implementation", "A lightweight Python project that is easy to run and extend.",
               "Streamlit UI · pandas data preparation · scikit-learn models · Plotly visualizations.")
    rect(slide, 0.7, 1.9, 5.85, 4.45, NAVY, rounded=True)
    textbox(slide, "PROJECT STRUCTURE", 0.98, 2.2, 4.9, 0.22, size=9, fill="AFC6DD", bold=True)
    tree = [
        ("student_analysis/", WHITE, True),
        ("├── app.py", "D9E6F2", False),
        ("├── student_analytics/", "D9E6F2", False),
        ("│   ├── data.py", "AFC6DD", False),
        ("│   └── analysis.py", "AFC6DD", False),
        ("├── data/student_performance.csv", "D9E6F2", False),
        ("├── scripts/", "D9E6F2", False),
        ("│   ├── generate_sample_data.py", "AFC6DD", False),
        ("│   └── create_project_presentation.py", "AFC6DD", False),
        ("├── tests/test_analysis.py", "D9E6F2", False),
        ("└── requirements.txt  ·  README.md", "D9E6F2", False),
    ]
    for i, (line, shade, bold) in enumerate(tree):
        textbox(slide, line, 1.0, 2.66 + i * 0.28, 5.05, 0.2, size=10, fill=shade, bold=bold)

    rect(slide, 6.88, 1.9, 5.75, 4.45, WHITE, rounded=True, line=LINE)
    textbox(slide, "RUN LOCALLY · WINDOWS POWERSHELL", 7.18, 2.2, 5.0, 0.22, size=9, fill=BLUE, bold=True)
    commands = [
        "py -3 -m venv .venv",
        ".\\.venv\\Scripts\\Activate.ps1",
        "pip install -r requirements.txt",
        "streamlit run app.py",
    ]
    for i, command in enumerate(commands):
        y = 2.8 + i * 0.62
        rect(slide, 7.18, y - 0.06, 5.12, 0.43, "F2F5F9", rounded=True)
        textbox(slide, command, 7.36, y + 0.04, 4.76, 0.18, size=10, fill=NAVY)
    textbox(slide, "Dashboard: http://localhost:8501", 7.18, 5.5, 4.8, 0.25, size=11, fill=GREEN, bold=True)
    textbox(slide, "Upload a CSV with the documented 13 required fields; filters and exports update with the selection.",
            7.18, 5.9, 4.95, 0.35, size=9.5, fill=MUTED)
    add_footer(slide, 8)

    # 9. Validation and responsible use
    slide = prs.slides.add_slide(blank)
    add_background(slide, NAVY)
    textbox(slide, "07 · VALIDATION & RESPONSIBLE USE", 0.7, 0.45, 11.8, 0.22, size=9, fill="AFC6DD", bold=True)
    textbox(slide, "Useful analysis starts with honest limits.", 0.7, 0.82, 11.8, 0.58, size=27, fill=WHITE, bold=True)
    rect(slide, 0.7, 1.82, 5.85, 3.55, "18395F", rounded=True)
    textbox(slide, "QUALITY CHECKS", 0.98, 2.12, 5.0, 0.22, size=9, fill="AFC6DD", bold=True)
    add_dot_bullet(slide, "5 unit tests pass: reproducibility, validation, classification, and clustering.",
                   1.0, 2.62, 5.1, accent=GREEN, size=11)
    add_dot_bullet(slide, "CSV schema, score ranges, numeric fields, and assessment dates are checked.",
                   1.0, 3.43, 5.1, accent=BLUE, size=11)
    add_dot_bullet(slide, "Risk inputs exclude final score and the derived target label.",
                   1.0, 4.24, 5.1, accent=ORANGE, size=11)
    rect(slide, 6.8, 1.82, 5.83, 3.55, "18395F", rounded=True)
    textbox(slide, "RESPONSIBLE DEPLOYMENT", 7.08, 2.12, 5.0, 0.22, size=9, fill="AFC6DD", bold=True)
    add_dot_bullet(slide, "Sample data is synthetic; performance does not establish real-world validity.",
                   7.1, 2.62, 5.05, accent=ORANGE, size=11)
    add_dot_bullet(slide, "Use as a support signal only; educators review context and make decisions.",
                   7.1, 3.43, 5.05, accent=TEAL, size=11)
    add_dot_bullet(slide, "Protect student privacy; obtain authorization and assess fairness on representative data.",
                   7.1, 4.24, 5.05, accent=RED, size=11)
    textbox(slide, "Next steps", 0.78, 5.85, 1.4, 0.28, size=13, fill=WHITE, bold=True)
    textbox(slide, "Evaluate on consented school data  →  measure subgroup outcomes  →  calibrate thresholds  →  review with educators",
            2.0, 5.86, 10.4, 0.32, size=11, fill="D9E6F2")
    add_footer(slide, 9, dark=True)

    prs.core_properties.title = "Student Academic Performance Analysis Using Data Mining"
    prs.core_properties.subject = "Streamlit dashboard, predictive analytics, and student engagement clustering"
    prs.core_properties.author = "Student Analytics Project"
    prs.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(f"Created {build_deck()}")
